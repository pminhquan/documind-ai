"""Export table results to JSON, CSV, and editable XLSX."""

import csv
import re
from pathlib import Path

from .models import Table, TableResult

# OCR text is untrusted input. CSV carries no cell-type metadata, so
# spreadsheet applications interpret leading "=", "+", "@" (and a
# leading "-" that is not a plain negative number) as formulas when the
# file opens. Such values get the standard quote-prefix escape. XLSX
# carries an explicit cell type instead (see export_xlsx).
CSV_ESCAPE_PREFIX = "'"
CSV_DANGEROUS_FIRST_CHARS = set("=+@")
NEGATIVE_NUMBER_RE = re.compile(r"-\d+(?:[.,]\d+)?")

# Control characters illegal in XLSX/XML strings (openpyxl would raise
# IllegalCharacterError); stripped from XLSX values only. JSON and CSV
# keep the raw recognized text.
XML_ILLEGAL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# Excel worksheet-title rules: these characters are forbidden, the title
# is capped at 31 characters, and titles must be unique case-insensitively
# within one workbook. Applied to SHEET NAMES only — JSON table_id values
# are never rewritten.
EXCEL_ILLEGAL_TITLE_CHARS_RE = re.compile(r"[\\/*?:\[\]]")
MAX_SHEET_TITLE_LENGTH = 31


def excel_sheet_title(table_id: str, used_titles: set[str]) -> str:
    """Deterministic, Excel-safe worksheet title for one table.

    Illegal characters become ``_``, whitespace-trimmed empties fall back
    to ``"Table"``, the result is capped at Excel's 31-character limit,
    and case-insensitive collisions get ``_2``, ``_3``… suffixes
    (truncation-aware so the title never exceeds the limit). No table is
    ever dropped; every table keeps a distinct sheet.
    """
    base = EXCEL_ILLEGAL_TITLE_CHARS_RE.sub("_", str(table_id)).strip()
    base = base[:MAX_SHEET_TITLE_LENGTH] or "Table"
    title, counter = base, 2
    while title.casefold() in used_titles:
        suffix = f"_{counter}"
        title = base[:MAX_SHEET_TITLE_LENGTH - len(suffix)] + suffix
        counter += 1
    used_titles.add(title.casefold())
    return title


def _csv_safe(value: str) -> str:
    """Quote-prefix values a spreadsheet would misread as formulas.

    The decision uses the first non-whitespace character; the prefix is
    inserted at position 0 so the original text (including leading
    whitespace) is preserved. A plain negative decimal (``-5.00``) is
    kept unchanged; any other leading ``-`` is escaped.
    """
    first = next((ch for ch in value if not ch.isspace()), "")
    if first in CSV_DANGEROUS_FIRST_CHARS:
        return CSV_ESCAPE_PREFIX + value
    if first == "-" and not NEGATIVE_NUMBER_RE.fullmatch(value.strip()):
        return CSV_ESCAPE_PREFIX + value
    return value


def _xlsx_text(value: str) -> str:
    """Strip XML-illegal control characters for XLSX cells."""
    return XML_ILLEGAL_CHARS_RE.sub("", value)


def table_matrix(table: Table) -> list[list[str]]:
    """Rebuild the rows x columns text grid from row-major cells."""
    matrix = [["" for _ in range(table.columns)] for _ in range(table.rows)]
    for cell in table.cells:
        matrix[cell.row_index][cell.column_index] = cell.text
    return matrix


def export_json(result: TableResult, path: str | Path, indent: int | None = 2) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(result.to_json(indent=indent), encoding="utf-8")
    return output


def _csv_paths(base: Path, table_count: int) -> list[Path]:
    if table_count == 0:
        return []
    if table_count == 1:
        return [base]
    return [base.with_name(f"{base.stem}_t{index}{base.suffix}")
            for index in range(table_count)]


def export_csv(result: TableResult, path: str | Path) -> list[Path]:
    """Write one CSV per table; multiple tables get ``_t<i>`` suffixes.

    Zero tables write no files (CSV is a per-table artifact).
    """
    base = Path(path)
    base.parent.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for table, target in zip(result.tables, _csv_paths(base, len(result.tables)), strict=True):
        with target.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerows(
                [_csv_safe(value) for value in row] for row in table_matrix(table)
            )
        written.append(target)
    return written


def export_xlsx(result: TableResult, path: str | Path) -> Path:
    """Write all tables into one workbook, one sheet per table.

    Zero tables produce a valid workbook with one empty default sheet.
    All values are stored with explicit string semantics (``data_type``
    ``"s"``), so OCR text can never become a formula cell.
    """
    from openpyxl import Workbook

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    if not result.tables:
        workbook.save(output)  # default "Sheet" stays: valid empty workbook
        return output

    workbook.remove(workbook.active)
    used_titles: set[str] = set()
    for table in result.tables:
        sheet = workbook.create_sheet(
            title=excel_sheet_title(table.table_id, used_titles))
        for row_index, row in enumerate(table_matrix(table), start=1):
            for col_index, value in enumerate(row, start=1):
                cell = sheet.cell(
                    row=row_index, column=col_index, value=_xlsx_text(value)
                )
                cell.data_type = "s"
    workbook.save(output)
    return output
