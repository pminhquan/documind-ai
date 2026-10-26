"""Web API tests — all inference mocked, no model downloads.

Covers: health, OCR/invoice/table/process endpoints (success + error
paths), upload security (size, extension, content), export (JSON/CSV/
XLSX with read-back), and filename sanitization.
"""

import csv
import io
import json
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image

from documind.web import service
from documind.web.app import create_app


@pytest.fixture
def client():
    return TestClient(create_app())


@pytest.fixture
def sample_png_bytes():
    buffer = io.BytesIO()
    Image.new("RGB", (32, 32), color="white").save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def mocked_engine(monkeypatch):
    """Replace the OCR engine with canned output; records predict calls."""

    class FakeEngine:
        def __init__(self):
            self.calls = 0

        def predict(self, path):
            self.calls += 1
            return [{
                # Shaped like a real receipt: supplier line first (no
                # amount), then TOTAL and its amount sharing one visual
                # row, so line grouping + keyword matching both work.
                "rec_texts": ["ACME STORE", "TOTAL", "13.53"],
                "rec_scores": [0.99, 0.99, 0.98],
                "rec_polys": [
                    [[1, 1], [60, 1], [60, 10], [1, 10]],
                    [[1, 20], [20, 20], [20, 29], [1, 29]],
                    [[25, 20], [40, 20], [40, 29], [25, 29]],
                ],
            }]

    engine = FakeEngine()
    monkeypatch.setattr(service, "_get_engine", lambda: engine)
    return engine


@pytest.fixture
def mocked_tables(monkeypatch):
    """Replace table recognition with a canned single table."""
    payload = {
        "source_image": "upload",
        "tables": [{
            "table_id": "table_0", "source_image": "upload", "rows": 1,
            "columns": 2,
            "cells": [
                {"row_index": 0, "column_index": 0, "text": "Item",
                 "bounding_box": [[0, 0], [1, 0], [1, 1], [0, 1]],
                 "source_ocr_indices": [0], "confidence": 0.9},
                {"row_index": 0, "column_index": 1, "text": "Qty",
                 "bounding_box": [[1, 0], [2, 0], [2, 1], [1, 1]],
                 "source_ocr_indices": [1], "confidence": 0.9},
            ],
            "warnings": [],
        }],
        "warnings": [],
    }
    monkeypatch.setattr(service, "tables_from_document",
                        lambda source, ocr_items=None: payload)
    return payload


def _upload(client, url, data=None, name="doc.png", content=None):
    content = content if content is not None else _png_bytes()
    with client as c:
        return c.post(url, files={"upload": (name, content)}, data=data or {})


def _png_bytes():
    buffer = io.BytesIO()
    Image.new("RGB", (16, 16), color="white").save(buffer, format="PNG")
    return buffer.getvalue()


class TestHealth:
    def test_health_reports_lazy_engine(self, client):
        response = client.get("/api/health")
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "ok"
        assert body["engine_loaded"] is False
        assert body["version"]

    def test_index_serves_html(self, client):
        response = client.get("/")
        assert response.status_code == 200
        assert "DocuMind AI" in response.text
        assert "app.js" in response.text


class TestOcrEndpoint:
    def test_success_with_mocked_engine(self, client, mocked_engine):
        response = _upload(client, "/api/ocr")
        assert response.status_code == 200
        body = response.json()
        assert body["items"][0]["text"] == "ACME STORE"
        assert body["items"][2]["text"] == "13.53"
        # full_text is one line per OCR region (M1 schema); the extraction
        # layer does its own line grouping.
        assert body["full_text"] == "ACME STORE\nTOTAL\n13.53"
        assert mocked_engine.calls == 1

    def test_engine_failure_maps_to_500(self, client, monkeypatch):
        def broken(path):
            raise service.DocumentProcessingError("OCR inference failed: boom")

        monkeypatch.setattr(service, "ocr_document", broken)
        response = _upload(client, "/api/ocr")
        assert response.status_code == 500
        assert "boom" in response.json()["detail"]

    def test_unsupported_extension_415(self, client):
        response = _upload(client, "/api/ocr", name="doc.gif", content=b"GIF89a")
        assert response.status_code == 415
        assert "Unsupported file type" in response.json()["detail"]

    def test_oversized_upload_413(self, client):
        response = _upload(client, "/api/ocr", name="big.png",
                           content=b"\0" * (10 * 1024 * 1024 + 1))
        assert response.status_code == 413

    def test_empty_upload_400(self, client):
        response = _upload(client, "/api/ocr", name="empty.png", content=b"")
        assert response.status_code == 400

    def test_non_image_content_400(self, client):
        response = _upload(client, "/api/ocr", name="fake.png",
                           content=b"definitely not an image")
        assert response.status_code == 400
        assert "decodable" in response.json()["detail"]

    def test_temp_files_cleaned_up(self, client, mocked_engine, tmp_path, monkeypatch):
        import tempfile

        temp_root = tmp_path / "tmp"
        temp_root.mkdir()
        monkeypatch.setattr(tempfile, "tempdir", str(temp_root))
        _upload(client, "/api/ocr")
        assert list(temp_root.iterdir()) == []


class TestInvoiceEndpoint:
    def test_success_runs_ocr_then_extraction(self, client, mocked_engine):
        response = _upload(client, "/api/invoice")
        assert response.status_code == 200
        body = response.json()
        assert body["invoice"]["validation_status"] in ("ok", "review_required")
        assert body["provenance"]["total"]["text"] == "13.53"

    def test_missing_fields_are_null(self, client, mocked_engine):
        body = _upload(client, "/api/invoice").json()["invoice"]
        assert body["supplier"] == "ACME STORE"  # first line: genuine extraction
        assert body["total"] == "13.53"
        assert body["discount"] is None     # absent from OCR text: null, never fabricated
        assert body["change_amount"] is None


class TestTableEndpoint:
    def test_success_with_canned_tables(self, client, mocked_tables):
        response = _upload(client, "/api/table")
        assert response.status_code == 200
        body = response.json()
        assert body["tables"][0]["rows"] == 1
        assert body["tables"][0]["cells"][0]["text"] == "Item"


class TestProcessEndpoint:
    def test_auto_runs_all_workflows(self, client, mocked_engine, mocked_tables):
        response = _upload(client, "/api/process", data={"workflow": "auto"})
        assert response.status_code == 200
        body = response.json()
        assert set(body["workflows"]) == {"ocr", "invoice", "table"}
        assert body["ocr"]["full_text"]
        assert body["invoice"]["invoice"]["total"] == "13.53"
        assert body["tables"]["tables"][0]["columns"] == 2

    def test_ocr_only_workflow(self, client, mocked_engine, mocked_tables):
        response = _upload(client, "/api/process", data={"workflow": "ocr"})
        body = response.json()
        assert body["ocr"] is not None
        assert body["invoice"] is None
        assert body["tables"] is None

    def test_unknown_workflow_422(self, client, mocked_engine):
        response = _upload(client, "/api/process", data={"workflow": "magic"})
        assert response.status_code == 422


class TestExport:
    @pytest.fixture
    def table_payload(self):
        return {
            "tables_result": {
                "source_image": "x.png",
                "tables": [{
                    "table_id": "table_0", "source_image": "x.png", "rows": 1,
                    "columns": 2,
                    "cells": [
                        {"row_index": 0, "column_index": 0, "text": "Item",
                         "bounding_box": [], "source_ocr_indices": []},
                        {"row_index": 0, "column_index": 1, "text": "=evil()",
                         "bounding_box": [], "source_ocr_indices": []},
                    ],
                    "warnings": [],
                }],
                "warnings": [],
            }
        }

    def test_json_export_preserves_machine_and_corrections(self, client, table_payload):
        response = client.post("/api/export", json={
            "kind": "table-json",
            "data": {"machine_output": table_payload["tables_result"],
                     "user_corrections": {"cells[0][0]": "Edited"}},
        })
        assert response.status_code == 200
        data = json.loads(response.content)
        assert data["machine_output"]["tables"][0]["cells"][0]["text"] == "Item"
        assert data["user_corrections"] == {"cells[0][0]": "Edited"}

    def test_csv_export_is_formula_safe(self, client, table_payload):
        import csv

        response = client.post("/api/export", json={
            "kind": "table-csv", "data": table_payload})
        assert response.status_code == 200
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8"))))
        assert rows[0] == ["Item", "'=evil()"]

    def test_xlsx_export_reads_back(self, client, table_payload):
        from openpyxl import load_workbook

        response = client.post("/api/export", json={
            "kind": "table-xlsx", "data": table_payload})
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook["table_0"]
        assert sheet.cell(row=1, column=1).value == "Item"
        assert sheet.cell(row=1, column=2).data_type != "f"

    def test_unknown_kind_422(self, client):
        response = client.post("/api/export", json={"kind": "exe", "data": {}})
        assert response.status_code == 422

    def test_non_object_data_422(self, client):
        response = client.post("/api/export", json={"kind": "ocr-json",
                                                    "data": "nope"})
        assert response.status_code == 422

    def test_filename_sanitized(self, client, table_payload):
        response = client.post("/api/export", json={
            "kind": "table-json", "data": table_payload,
            "filename": "../../evil name.json"})
        disposition = response.headers["Content-Disposition"]
        name = disposition.split('filename="')[1].rstrip('"')
        assert "/" not in name and "\\" not in name and " " not in name


def _two_table_payload():
    def table(index):
        return {
            "table_id": f"table_{index}", "source_image": "x.png", "rows": 1,
            "columns": 2,
            "cells": [
                {"row_index": 0, "column_index": 0, "text": f"A{index}",
                 "bounding_box": [], "source_ocr_indices": []},
                {"row_index": 0, "column_index": 1, "text": f"B{index}",
                 "bounding_box": [], "source_ocr_indices": []},
            ],
            "warnings": [],
        }

    return {"source_image": "x.png", "tables": [table(0), table(1)], "warnings": []}


class TestCorrectionMaterialization:
    """A1: table corrections must reach CSV/XLSX exports via a validated
    materialization step; JSON keeps machine output + corrections apart."""

    @staticmethod
    def _hitl_payload(corrections):
        machine = {
            "source_image": "x.png",
            "tables": [{
                "table_id": "table_0", "source_image": "x.png", "rows": 2,
                "columns": 2,
                "cells": [
                    {"row_index": 0, "column_index": 0, "text": "Item",
                     "bounding_box": [], "source_ocr_indices": []},
                    {"row_index": 0, "column_index": 1, "text": "Qty",
                     "bounding_box": [], "source_ocr_indices": []},
                    {"row_index": 1, "column_index": 0, "text": "Coffee",
                     "bounding_box": [], "source_ocr_indices": []},
                    {"row_index": 1, "column_index": 1, "text": "3.50",
                     "bounding_box": [], "source_ocr_indices": []},
                ],
                "warnings": [],
            }],
            "warnings": [],
        }
        return {"machine_output": machine, "user_corrections": corrections}

    def test_csv_uses_corrected_values(self, client):
        response = client.post("/api/export", json={
            "kind": "table-csv",
            "data": self._hitl_payload({"tables[0].cells[1][1]": "4.00"}),
        })
        assert response.status_code == 200
        import csv
        import io

        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8"))))
        assert rows[1] == ["Coffee", "4.00"]

    def test_xlsx_uses_corrected_values(self, client):
        from openpyxl import load_workbook

        response = client.post("/api/export", json={
            "kind": "table-xlsx",
            "data": self._hitl_payload({"tables[0].cells[1][1]": "4.00"}),
        })
        sheet = load_workbook(io.BytesIO(response.content))["table_0"]
        assert sheet.cell(row=2, column=2).value == "4.00"

    def test_unedited_values_unchanged_and_machine_untouched(self, client):
        payload = self._hitl_payload({"tables[0].cells[1][1]": "4.00"})
        response = client.post("/api/export", json={
            "kind": "table-json", "data": payload})
        data = json.loads(response.content)
        assert data["machine_output"]["tables"][0]["cells"][3]["text"] == "3.50"
        assert data["user_corrections"] == {"tables[0].cells[1][1]": "4.00"}

    def test_out_of_bounds_correction_422(self, client):
        response = client.post("/api/export", json={
            "kind": "table-csv",
            "data": self._hitl_payload({"tables[0].cells[9][9]": "x"}),
        })
        assert response.status_code == 422
        assert "out of bounds" in response.json()["detail"]

    def test_foreign_path_correction_422(self, client):
        response = client.post("/api/export", json={
            "kind": "table-csv",
            "data": self._hitl_payload({"invoice.total": "1.00"}),
        })
        assert response.status_code == 422
        assert "correction path" in response.json()["detail"]


class TestMultiTableExport:
    """A2: multiple tables -> ZIP of per-table CSVs; one table -> CSV."""

    def test_single_table_returns_csv(self, client):
        single = {"source_image": "x.png",
                  "tables": _two_table_payload()["tables"][:1], "warnings": []}
        response = client.post("/api/export", json={
            "kind": "table-csv", "data": {"machine_output": single}})
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert response.headers["Content-Disposition"].endswith('tables.csv"')

    def test_multiple_tables_return_zip(self, client):
        response = client.post("/api/export", json={
            "kind": "table-csv", "data": {"machine_output": _two_table_payload()}})
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/zip"
        assert response.headers["Content-Disposition"].endswith('tables.zip"')
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            names = sorted(archive.namelist())
            assert names == ["tables_t0.csv", "tables_t1.csv"]
            rows = list(csv.reader(
                io.StringIO(archive.read("tables_t0.csv").decode("utf-8"))))
            assert rows[0] == ["A0", "B0"]

    def test_zero_tables_422(self, client):
        response = client.post("/api/export", json={
            "kind": "table-csv",
            "data": {"machine_output": {"source_image": "x", "tables": [],
                                        "warnings": []}}})
        assert response.status_code == 422


class TestInferenceSafety:
    """A3: one shared engine, serialized predict(), no duplicate engines,
    no redundant OCR in unified processing."""

    @pytest.fixture
    def counting_engine(self, monkeypatch):
        import threading
        import time

        from documind.web import service

        class CountingEngine:
            def __init__(self):
                self.calls = 0
                self.max_concurrent = 0
                self.active = 0
                self._counter_lock = threading.Lock()
                self.fail_next = False

            def predict(self, path):
                with self._counter_lock:
                    self.calls += 1
                    self.active += 1
                    self.max_concurrent = max(self.max_concurrent, self.active)
                try:
                    time.sleep(0.05)
                    if self.fail_next:
                        self.fail_next = False
                        raise RuntimeError("transient inference failure")
                    return [{
                        "rec_texts": ["ACME STORE", "TOTAL", "13.53"],
                        "rec_scores": [0.99, 0.99, 0.98],
                        "rec_polys": [
                            [[1, 1], [60, 1], [60, 10], [1, 10]],
                            [[1, 20], [20, 20], [20, 29], [1, 29]],
                            [[25, 20], [40, 20], [40, 29], [25, 29]],
                        ],
                    }]
                finally:
                    with self._counter_lock:
                        self.active -= 1

        engine = CountingEngine()
        monkeypatch.setattr(service, "_get_engine", lambda: engine)
        return engine

    def test_concurrent_ocr_serialized(self, counting_engine):
        from concurrent.futures import ThreadPoolExecutor

        from documind.web import service

        source = Path("whatever.png")  # engine is faked; path unused
        with ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(lambda _: service.ocr_document(source), range(4)))
        assert counting_engine.calls == 4
        assert counting_engine.max_concurrent == 1

    def test_unified_auto_avoids_redundant_ocr(self, counting_engine, tmp_path):
        from documind.web import service

        image = tmp_path / "doc.png"
        Image.new("RGB", (40, 40), color="white").save(image)
        service.process_unified(image, ["ocr", "invoice", "table"])
        assert counting_engine.calls == 1  # tables reuse the OCR items

    def test_table_only_uses_shared_engine_runner(self, counting_engine, tmp_path,
                                                  monkeypatch):
        from documind.web import service

        constructed = []

        def forbidden_engine(**kwargs):
            constructed.append(kwargs)
            raise AssertionError("table pipeline must not build its own engine")

        # the pipeline imports the engine lazily from its source module
        monkeypatch.setattr("documind.ocr.engine.PaddleOcrEngine", forbidden_engine)
        # A ruled grid is required so table detection proceeds to OCR.
        import cv2
        import numpy as np

        image_array = np.full((120, 220, 3), 255, dtype=np.uint8)
        for y in (10, 60, 110):
            cv2.line(image_array, (10, y), (210, y), (0, 0, 0), 2)
        for x in (10, 110, 210):
            cv2.line(image_array, (x, 10), (x, 110), (0, 0, 0), 2)
        image = tmp_path / "doc.png"
        cv2.imwrite(str(image), image_array)
        service.tables_from_document(image)
        assert counting_engine.calls == 1
        assert constructed == []

    def test_failed_inference_does_not_poison_engine(self, counting_engine):
        from documind.web import service

        counting_engine.fail_next = True
        with pytest.raises(service.DocumentProcessingError):
            service.ocr_document(Path("whatever.png"))
        result = service.ocr_document(Path("whatever.png"))
        assert result["items"][0]["text"] == "ACME STORE"
        assert counting_engine.calls == 2

    def test_reset_engine_is_safe(self, monkeypatch):
        from documind.web import service

        service.reset_engine()
        assert service._engine is None


def _valid_export_table(table_id: str = "table_0") -> dict:
    return {
        "table_id": table_id, "source_image": "x.png", "rows": 2, "columns": 2,
        "cells": [
            {"row_index": r, "column_index": c, "text": f"r{r}c{c}",
             "bounding_box": [[0, 0], [1, 0], [1, 1], [0, 1]],
             "source_ocr_indices": [], "confidence": 0.9}
            for r in range(2) for c in range(2)
        ],
        "warnings": [],
    }


def _tables_payload(table: dict, **collection_overrides) -> dict:
    """Wrap one table in the TableResponse collection shape."""
    return {"tables": [table], "warnings": [], **collection_overrides}


def _mutate_first_cell(table: dict, **changes) -> dict:
    mutated = json.loads(json.dumps(table))
    mutated["cells"][0].update(changes)
    return mutated


# Each entry mutates the TableResponse COLLECTION; all are client errors
# that must come back as 422 instead of an uncontrolled 500.
MALFORMED_TABLE_PAYLOADS = {
    "tables_not_list": lambda t: {"tables": "junk"},
    "non_object_table": lambda t: {"tables": ["junk"]},
    "rows_string": lambda t: _tables_payload({**t, "rows": "abc"}),
    "rows_bool": lambda t: _tables_payload({**t, "rows": True}),
    "rows_negative": lambda t: _tables_payload({**t, "rows": -2}),
    "rows_too_large": lambda t: _tables_payload({**t, "rows": 501}),
    "columns_bool": lambda t: _tables_payload({**t, "columns": False}),
    "too_many_cells": lambda t: _tables_payload({**t, "rows": 200,
                                                 "columns": 200}),
    "cells_not_list": lambda t: _tables_payload({**t, "cells": "junk"}),
    "cell_not_object": lambda t: _tables_payload({**t, "cells": ["junk"]}),
    "cell_missing_bounding_box": lambda t: _tables_payload({
        **t, "cells": [{k: v for k, v in t["cells"][0].items()
                        if k != "bounding_box"}, *t["cells"][1:]]}),
    "cell_unknown_key": lambda t: _tables_payload(
        _mutate_first_cell(t, junk_key=1)),
    "cell_coord_string": lambda t: _tables_payload(
        _mutate_first_cell(t, row_index="0")),
    "cell_coord_bool": lambda t: _tables_payload(
        _mutate_first_cell(t, row_index=True)),
    "cell_coord_out_of_bounds": lambda t: _tables_payload(
        _mutate_first_cell(t, row_index=5)),
    "cell_coord_negative": lambda t: _tables_payload(
        _mutate_first_cell(t, row_index=-1)),
    "cell_text_not_string": lambda t: _tables_payload(
        _mutate_first_cell(t, text=42)),
    "warnings_not_list": lambda t: _tables_payload(t, warnings="junk"),
}


class TestTableExportPayloadValidation:
    """D-17: malformed table export payloads must return 422, never an
    uncontrolled 500 from int()/TableCell(**cell)/openpyxl."""

    @pytest.fixture
    def valid_table(self):
        return _valid_export_table()

    @pytest.mark.parametrize("case", sorted(MALFORMED_TABLE_PAYLOADS))
    def test_malformed_payload_is_422_not_500(self, client, valid_table, case):
        payload = {"machine_output": MALFORMED_TABLE_PAYLOADS[case](valid_table)}
        response = client.post("/api/export", json={
            "kind": "table-csv", "data": payload})
        assert response.status_code == 422, (
            f"case {case}: expected 422, got {response.status_code}")

    def test_valid_table_csv_still_exports(self, client, valid_table):
        response = client.post("/api/export", json={
            "kind": "table-csv",
            "data": {"machine_output": _tables_payload(valid_table)}})
        assert response.status_code == 200
        assert "r1c1" in response.content.decode("utf-8")

    def test_valid_xlsx_read_back_preserves_cells(self, client, valid_table):
        import openpyxl
        response = client.post("/api/export", json={
            "kind": "table-xlsx",
            "data": {"machine_output": _tables_payload(valid_table)}})
        assert response.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
        assert ws["B2"].value == "r1c1"

    def test_formula_injection_safety_intact(self, client, valid_table):
        payload = _tables_payload(_mutate_first_cell(valid_table, text="=cmd()"))
        response = client.post("/api/export", json={
            "kind": "table-csv", "data": {"machine_output": payload}})
        assert response.status_code == 200
        assert "'=cmd()" in response.content.decode("utf-8")

    def test_xlsx_sheet_title_sanitized(self, client, valid_table):
        import openpyxl
        illegal = "\\/*?:[]"
        payload = _tables_payload({**valid_table, "table_id": "bad:title/here*?"})
        response = client.post("/api/export", json={
            "kind": "table-xlsx", "data": {"machine_output": payload}})
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        (title,) = workbook.sheetnames
        assert not (set(title) & set(illegal))
        assert workbook[title]["B2"].value == "r1c1"

    def test_xlsx_sheet_title_truncated_to_31(self, client, valid_table):
        import openpyxl
        payload = _tables_payload({**valid_table, "table_id": "t" * 40})
        response = client.post("/api/export", json={
            "kind": "table-xlsx", "data": {"machine_output": payload}})
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        (title,) = workbook.sheetnames
        assert len(title) <= 31
        assert workbook[title]["B2"].value == "r1c1"

    def test_xlsx_duplicate_titles_are_unique(self, client, valid_table):
        import openpyxl
        payload = {"tables": [_valid_export_table("Report"),
                              _valid_export_table("report")], "warnings": []}
        response = client.post("/api/export", json={
            "kind": "table-xlsx", "data": {"machine_output": payload}})
        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        assert len(workbook.sheetnames) == 2
        assert len({name.casefold() for name in workbook.sheetnames}) == 2

    def test_batch_zip_inner_malformed_table_is_422(self, client, valid_table):
        payload = {"files": [{
            "kind": "table-xlsx", "filename": "t.xlsx",
            "data": {"machine_output": _tables_payload({**valid_table,
                                                        "rows": "abc"})},
        }]}
        response = client.post("/api/export", json={
            "kind": "batch-zip", "data": payload})
        assert response.status_code == 422

class TestResponseSourceSanitization:
    """D-19: no server-local temporary path may appear in any response."""

    TEMP_MARKER = "documind_upload_"

    def _uploaded(self, client, url, data=None):
        return client.post(url, files={"upload": ("doc.png", _png_bytes())},
                           data=data or {})

    def test_ocr_response_has_no_temp_path(self, client, mocked_engine):
        response = self._uploaded(client, "/api/ocr")
        assert response.status_code == 200
        assert self.TEMP_MARKER not in response.text
        assert response.json()["source"] == "upload"

    def test_invoice_response_has_no_temp_path(self, client, mocked_engine):
        response = self._uploaded(client, "/api/invoice")
        assert response.status_code == 200
        assert self.TEMP_MARKER not in response.text
        assert response.json()["source"] == "upload"

    def test_table_response_has_no_temp_path(self, client, mocked_engine,
                                             mocked_tables):
        response = self._uploaded(client, "/api/table")
        assert response.status_code == 200
        assert self.TEMP_MARKER not in response.text
        assert response.json()["source_image"] == "upload"

    def test_unified_response_has_no_temp_path(self, client, mocked_engine,
                                               mocked_tables):
        response = self._uploaded(client, "/api/process",
                                  data={"workflow": "auto"})
        assert response.status_code == 200
        assert self.TEMP_MARKER not in response.text
        body = response.json()
        assert body["source"] == "upload"
        assert body["ocr"]["source"] == "upload"
        assert body["tables"]["source_image"] == "upload"
        assert body["invoice"]["source"] == "upload"

    def test_engine_error_message_has_no_temp_path(self, client, monkeypatch):
        from documind.ocr.engine import OcrEngineError

        class LeakyEngine:
            def predict(self, path):
                raise OcrEngineError(f"OCR inference failed for '{path}': boom")

        monkeypatch.setattr(service, "_get_engine", lambda: LeakyEngine())
        response = self._uploaded(client, "/api/ocr")
        assert response.status_code == 500
        assert self.TEMP_MARKER not in response.text
        assert "boom" in response.text  # useful context survives

    def test_decode_error_message_has_no_temp_path(self, client, mocked_engine,
                                                   monkeypatch):
        from documind.tables.pipeline import TableRecognitionError

        def leaky_recognize(source, ocr_items=None, ocr_runner=None):
            raise TableRecognitionError(f"Cannot decode image '{source}'")

        monkeypatch.setattr(service, "recognize_tables", leaky_recognize)
        response = self._uploaded(client, "/api/table")
        assert response.status_code == 500
        assert self.TEMP_MARKER not in response.text
        assert "Cannot decode image" in response.text
