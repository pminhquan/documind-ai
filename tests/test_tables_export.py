"""JSON / CSV / XLSX export tests."""

import csv
import json

import pytest

from documind.tables.exporters import (
    export_csv,
    export_json,
    export_xlsx,
    table_matrix,
)
from documind.tables.models import Table, TableCell, TableResult


def make_result(tables=1):
    def cell(r, c, text):
        return TableCell(
            row_index=r, column_index=c, text=text,
            bounding_box=[[c * 10, r * 10], [c * 10 + 10, r * 10],
                          [c * 10 + 10, r * 10 + 10], [c * 10, r * 10 + 10]],
            source_ocr_indices=[r * 2 + c] if text else [],
            confidence=0.9 if text else None,
        )

    def build(index):
        return Table(
            table_id=f"table_{index}",
            source_image="img.png",
            rows=2,
            columns=2,
            cells=[cell(0, 0, "Item"), cell(0, 1, "Qty"),
                   cell(1, 0, "Coffee"), cell(1, 1, "")],
            warnings=[],
        )

    return TableResult(
        source_image="img.png",
        tables=[build(i) for i in range(tables)],
        warnings=[],
    )


class TestZeroTableExport:
    """Regression: export_xlsx removed the only sheet and saved zero
    sheets -> openpyxl 'At least one sheet must be visible' crash."""

    def test_xlsx_with_zero_tables_is_a_valid_empty_workbook(self, tmp_path):
        from openpyxl import load_workbook

        result = TableResult(source_image="blank.png", tables=[], warnings=[])
        path = export_xlsx(result, tmp_path / "out.xlsx")
        workbook = load_workbook(path)
        assert workbook.sheetnames == ["Sheet"]

    def test_csv_with_zero_tables_writes_no_files(self, tmp_path):
        result = TableResult(source_image="blank.png", tables=[], warnings=[])
        written = export_csv(result, tmp_path / "out.csv")
        assert written == []
        assert not (tmp_path / "out.csv").exists()


class TestFormulaInjection:
    """Regression: OCR text is untrusted; exported cells must never
    become executable spreadsheet formulas."""

    @pytest.fixture
    def dangerous_result(self):
        def make(column_index, text):
            return TableCell(
                row_index=0, column_index=column_index, text=text,
                bounding_box=[[column_index, 0], [column_index + 1, 0],
                              [column_index + 1, 1], [column_index, 1]],
            )

        table = Table(
            table_id="table_0", source_image="img.png", rows=1, columns=6,
            cells=[
                make(0, "=cmd()"),
                make(1, "+sum(A1)"),
                make(2, "@import"),
                make(3, "-5.00"),        # legitimate negative number: keep
                make(4, "-notanumber"),  # '-' but not numeric: escape
                make(5, " \x01=hidden()"),  # leading whitespace + control char
            ],
        )
        return TableResult(source_image="img.png", tables=[table], warnings=[])

    def test_xlsx_stores_strings_not_formulas(self, dangerous_result, tmp_path):
        from openpyxl import load_workbook

        path = export_xlsx(dangerous_result, tmp_path / "out.xlsx")
        sheet = load_workbook(path)["table_0"]
        values = [sheet.cell(row=1, column=c).value for c in range(1, 7)]
        assert values == [
            "=cmd()", "+sum(A1)", "@import", "-5.00", "-notanumber", " =hidden()",
        ]
        for c in range(1, 7):
            assert sheet.cell(row=1, column=c).data_type != "f"
        assert sheet.cell(row=1, column=1).data_type == "s"

    def test_csv_uses_quote_prefix_policy(self, dangerous_result, tmp_path):
        written = export_csv(dangerous_result, tmp_path / "out.csv")
        with written[0].open(encoding="utf-8", newline="") as handle:
            rows = list(csv.reader(handle))
        assert rows[0] == [
            "'=cmd()", "'+sum(A1)", "'@import", "-5.00", "'-notanumber",
            " \x01=hidden()",
        ]

    def test_json_keeps_raw_text_unchanged(self, dangerous_result, tmp_path):
        path = export_json(dangerous_result, tmp_path / "out.json")
        data = json.loads(path.read_text(encoding="utf-8"))
        texts = [c["text"] for c in data["tables"][0]["cells"]]
        assert texts == [
            "=cmd()", "+sum(A1)", "@import", "-5.00", "-notanumber", " \x01=hidden()",
        ]

    def test_benign_numeric_text_is_untouched(self, tmp_path):
        table = Table(
            table_id="table_0", source_image="img.png", rows=1, columns=2,
            cells=[
                TableCell(row_index=0, column_index=0, text="3.50",
                          bounding_box=[[0, 0], [1, 0], [1, 1], [0, 1]]),
                TableCell(row_index=0, column_index=1, text="Coffee",
                          bounding_box=[[1, 0], [2, 0], [2, 1], [1, 1]]),
            ],
        )
        written = export_csv(
            TableResult(source_image="img.png", tables=[table]), tmp_path / "o.csv"
        )
        with written[0].open(encoding="utf-8", newline="") as handle:
            assert list(csv.reader(handle)) == [["3.50", "Coffee"]]


class TestTableMatrix:
    def test_matrix_preserves_empty_cells(self):
        assert table_matrix(make_result().tables[0]) == [
            ["Item", "Qty"],
            ["Coffee", ""],
        ]


class TestJsonExport:
    def test_round_trip(self, tmp_path):
        path = export_json(make_result(), tmp_path / "out.json")
        data = json.loads(path.read_text(encoding="utf-8"))
        assert data["source_image"] == "img.png"
        assert data["tables"][0]["rows"] == 2
        cell = data["tables"][0]["cells"][0]
        assert cell["text"] == "Item"
        assert cell["row_index"] == 0 and cell["column_index"] == 0
        assert len(cell["bounding_box"]) == 4
        assert cell["confidence"] == 0.9
        empty = data["tables"][0]["cells"][3]
        assert empty["text"] == "" and empty["confidence"] is None


class TestCsvExport:
    def test_single_table_uses_exact_path(self, tmp_path):
        written = export_csv(make_result(), tmp_path / "out.csv")
        assert written == [tmp_path / "out.csv"]
        with written[0].open(encoding="utf-8", newline="") as handle:
            rows = list(csv.reader(handle))
        assert rows == [["Item", "Qty"], ["Coffee", ""]]

    def test_multiple_tables_get_suffixes(self, tmp_path):
        written = export_csv(make_result(tables=2), tmp_path / "out.csv")
        assert [p.name for p in written] == ["out_t0.csv", "out_t1.csv"]
        with written[0].open(encoding="utf-8", newline="") as handle:
            assert list(csv.reader(handle)) == [["Item", "Qty"], ["Coffee", ""]]


class TestXlsxExport:
    def test_workbook_is_readable_and_editable(self, tmp_path):
        from openpyxl import load_workbook

        path = export_xlsx(make_result(tables=2), tmp_path / "out.xlsx")
        workbook = load_workbook(path)
        assert workbook.sheetnames == ["table_0", "table_1"]
        sheet = workbook["table_0"]
        assert sheet.cell(row=1, column=1).value == "Item"
        assert sheet.cell(row=1, column=2).value == "Qty"
        assert sheet.cell(row=2, column=1).value == "Coffee"
        assert sheet.cell(row=2, column=2).value in ("", None)


class TestExcelSheetTitleSafety:
    """openpyxl silently SAVES worksheet titles containing XML-illegal
    control characters or unpaired surrogates, but the produced
    xl/workbook.xml is not well-formed XML 1.0 — Excel and every strict
    parser reject the whole file. excel_sheet_title() must apply the
    same character policy to titles that _xlsx_text() applies to cell
    values."""

    @staticmethod
    def _single_table(table_id):
        return Table(
            table_id=table_id, source_image="img.png", rows=1, columns=1,
            cells=[TableCell(row_index=0, column_index=0, text="x",
                             bounding_box=[[0, 0], [1, 0], [1, 1], [0, 1]],
                             source_ocr_indices=[], confidence=0.9)],
            warnings=[],
        )

    @staticmethod
    def _workbook_xml_bytes(path):
        import zipfile
        with zipfile.ZipFile(path) as archive:
            return archive.read("xl/workbook.xml")

    def _assert_valid_artifact(self, tmp_path, table_id):
        import xml.etree.ElementTree as ET

        from openpyxl import load_workbook

        path = export_xlsx(TableResult(source_image="img.png",
                                       tables=[self._single_table(table_id)]),
                           tmp_path / "out.xlsx")
        raw = self._workbook_xml_bytes(path)
        ET.fromstring(raw)  # strict XML 1.0: must be well-formed
        workbook = load_workbook(path)
        assert workbook.sheetnames, "no table may be dropped"
        return workbook.sheetnames[0]

    @pytest.mark.parametrize("table_id,expected", [
        ("bad\x00title", "badtitle"),
        ("bad\x01title", "badtitle"),
        ("bad\x0btitle", "badtitle"),
        ("bad\ud800title", "badtitle"),
    ])
    def test_xml_illegal_title_characters_are_stripped(
            self, tmp_path, table_id, expected):
        title = self._assert_valid_artifact(tmp_path, table_id)
        assert title == expected

    def test_control_only_title_falls_back_to_table(self, tmp_path):
        title = self._assert_valid_artifact(tmp_path, "\x00\x01\x0b")
        assert title == "Table"

    def test_valid_vietnamese_title_preserved(self, tmp_path):
        title = self._assert_valid_artifact(tmp_path, "Bảng Doanh Thu 2026")
        assert title == "Bảng Doanh Thu 2026"

    def test_normal_titles_unchanged_after_fix(self, tmp_path):
        title = self._assert_valid_artifact(tmp_path, "table_0")
        assert title == "table_0"
