"""CLI tests for documind-table (mocked OCR via --ocr-json)."""

import json
from pathlib import Path

import cv2
import numpy as np
import pytest
from PIL import Image

from documind import table_cli
from documind.tables.geometry import detect_table_grids

SAMPLE = Path(__file__).parent.parent / "assets" / "samples" / "table_products.png"
TRUTH = (
    Path(__file__).parent / "fixtures" / "tables" / "table_products.truth.json"
)


def build_ocr_json_from_truth(image_path: Path, truth: dict, output: Path) -> None:
    """Craft an OCR JSON whose regions sit at the truth cell centers.

    Uses real grid geometry so the fixture exercises the same assignment
    path as genuine OCR output - without running the model.
    """
    image = cv2.imread(str(image_path))
    grid = detect_table_grids(image)[0]
    items = []
    for row_index, row in enumerate(truth["cells"]):
        for col_index, text in enumerate(row):
            if not text:
                continue
            rect = grid.cell_rect(row_index, col_index)
            cx = rect.x + rect.w // 2
            cy = rect.y + rect.h // 2
            items.append({
                "text": text,
                "confidence": 0.99,
                "box": [[cx - 20, cy - 8], [cx + 20, cy - 8],
                        [cx + 20, cy + 8], [cx - 20, cy + 8]],
            })
    output.write_text(
        json.dumps({"engine": "fixture", "items": items}), encoding="utf-8"
    )


class TestTableCliHelp:
    def test_help_exits_zero(self, capsys):
        with pytest.raises(SystemExit) as exc_info:
            table_cli.main(["--help"])
        assert exc_info.value.code == 0
        assert "documind-table" in capsys.readouterr().out


class TestTableCliErrors:
    def test_missing_input_returns_2(self, capsys):
        assert table_cli.main(["missing_table.png"]) == 2
        assert "not found" in capsys.readouterr().err

    def test_undecodable_image_returns_1(self, tmp_path, capsys):
        fake = tmp_path / "fake.png"
        fake.write_bytes(b"not really a png image")
        assert table_cli.main([str(fake)]) == 1
        assert "Table recognition error" in capsys.readouterr().err

    def test_unsupported_extension_returns_2(self, tmp_path, capsys):
        path = tmp_path / "doc.gif"
        path.write_bytes(b"x")
        assert table_cli.main([str(path)]) == 2
        assert "Unsupported" in capsys.readouterr().err


class TestTableCliNoTable:
    def test_blank_image_reports_warning_and_exits_zero(self, tmp_path):
        blank = tmp_path / "blank.png"
        Image.new("RGB", (200, 200), color="white").save(blank)
        output = tmp_path / "out.json"
        code = table_cli.main([str(blank), "-o", str(output)])
        assert code == 0
        data = json.loads(output.read_text(encoding="utf-8"))
        assert data["tables"] == []
        assert any("no ruled table" in w for w in data["warnings"])


class TestTableCliHappyPath:
    @pytest.fixture
    def run_args(self, tmp_path):
        truth = json.loads(TRUTH.read_text(encoding="utf-8"))
        ocr_json = tmp_path / "ocr.json"
        build_ocr_json_from_truth(SAMPLE, truth, ocr_json)
        output = tmp_path / "tables.json"
        csv_path = tmp_path / "tables.csv"
        xlsx_path = tmp_path / "tables.xlsx"
        return truth, ocr_json, output, csv_path, xlsx_path

    def test_full_pipeline_matches_ground_truth(
        self, run_args
    ):
        truth, ocr_json, output, csv_path, xlsx_path = run_args
        code = table_cli.main([
            str(SAMPLE), "-o", str(output),
            "--csv", str(csv_path), "--xlsx", str(xlsx_path),
            "--ocr-json", str(ocr_json),
        ])
        assert code == 0

        data = json.loads(output.read_text(encoding="utf-8"))
        assert len(data["tables"]) == 1
        table = data["tables"][0]
        assert table["rows"] == truth["rows"]
        assert table["columns"] == truth["columns"]

        matrix = [["" for _ in range(table["columns"])] for _ in range(table["rows"])]
        for cell in table["cells"]:
            matrix[cell["row_index"]][cell["column_index"]] = cell["text"]
        assert matrix == truth["cells"]

    def test_csv_and_xlsx_written(self, run_args):
        _, ocr_json, output, csv_path, xlsx_path = run_args
        code = table_cli.main([
            str(SAMPLE), "-o", str(output),
            "--csv", str(csv_path), "--xlsx", str(xlsx_path),
            "--ocr-json", str(ocr_json),
        ])
        assert code == 0
        assert csv_path.exists()
        assert xlsx_path.exists()
        from openpyxl import load_workbook

        sheet = load_workbook(xlsx_path)["table_0"]
        assert sheet.cell(row=1, column=1).value == "Item"

    def test_default_output_path(self, tmp_path, monkeypatch):
        monkeypatch.chdir(tmp_path)
        truth = json.loads(TRUTH.read_text(encoding="utf-8"))
        ocr_json = tmp_path / "ocr.json"
        build_ocr_json_from_truth(SAMPLE, truth, ocr_json)
        code = table_cli.main([str(SAMPLE), "--ocr-json", str(ocr_json)])
        assert code == 0
        assert (tmp_path / "outputs" / "table_products.tables.json").exists()
